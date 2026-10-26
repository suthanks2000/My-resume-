# -*- coding: utf-8 -*-
"""Build Suthan K's resume as .docx (python-docx) and .pdf (ReportLab).

Both writers follow the same fixed template from resume_spec.md:
A4, single column, 0.6 in margins, Arial, name 18 pt bold, title line 10.5 pt bold
navy, section headings 11 pt bold CAPS with a thin rule, body 10 pt.

LibreOffice is not installed on this box (Amazon Linux 2023, arm64), so the PDF is
rendered directly instead of converted from the .docx.

Usage:
    python3.13 build_resume.py <outdir> [role ...]
"""
import glob
import os
import sys
from datetime import datetime

import resume_data as D

# ---------------------------------------------------------------- shared bits
NAVY_HEX = "002060"
NAVY_RGB = (0x00 / 255, 0x20 / 255, 0x60 / 255)
RULE_HEX = "002060"

PT = 1.0
A4_W_PT, A4_H_PT = 595.276, 841.890  # 210 x 297 mm
MARGIN_PT = 0.6 * 72

# Vertical rhythm, in points. Tuned so the full content lands on one page.
BODY_SIZE = 10
BODY_LEAD = 11.6
BULLET_AFTER = 1.0
SECTION_BEFORE = 5.0
SECTION_AFTER = 2.5
JOB_AFTER = 0.8
BLOCK_AFTER = 3.0
BULLET_INDENT_PT = 10.5


def role_items(roles):
    for key in roles:
        if key not in D.ROLES:
            raise SystemExit(f"unknown role: {key} (known: {', '.join(D.ROLES)})")
        yield key, D.ROLES[key]


def out_name(suffix, ext):
    stem = "Suthan_K_Resume" if not suffix else f"Suthan_K_Resume_{suffix}"
    return f"{stem}.{ext}"


# ------------------------------------------------------------------- the .docx
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.shared import Mm, Pt, Inches, RGBColor  # noqa: E402

# Schema order of w:pPr children: anything we inject by hand must respect it or
# Word complains the file is corrupt.
PPR_SEQ = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]


def _localname(el):
    return el.tag.split("}")[-1]


def _reorder_ppr(pPr):
    rank = {name: i for i, name in enumerate(PPR_SEQ)}
    for el in sorted(pPr, key=lambda e: rank.get(_localname(e), 999)):
        pPr.append(el)


def _add_bottom_rule(par):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")  # 6/8 pt = 0.75 pt, a thin line
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RULE_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)
    _reorder_ppr(pPr)


def _para(doc, *, size=BODY_SIZE, bold=False, color=None, align=None,
          before=0.0, after=BULLET_AFTER, indent=None, hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    pf.widow_control = False
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Pt(indent)
    if hanging is not None:
        pf.first_line_indent = Pt(-hanging)
    return p


def _run(p, text, *, size=BODY_SIZE, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    # Make the font stick for every script range, not just Latin.
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Arial")
    return r


def _section_heading(doc, text):
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT,
              before=SECTION_BEFORE, after=SECTION_AFTER)
    _run(p, text.upper(), size=11, bold=True)
    _add_bottom_rule(p)
    return p


def _bullet(doc, text, *, after=BULLET_AFTER):
    p = _para(doc, indent=BULLET_INDENT_PT, hanging=BULLET_INDENT_PT, after=after)
    _run(p, "•  " + text)
    return p


def build_docx(path, title_line, summary, skill_order):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(BODY_SIZE)
    nrPr = normal.element.get_or_add_rPr()
    nrFonts = nrPr.find(qn("w:rFonts"))
    if nrFonts is None:
        nrFonts = OxmlElement("w:rFonts")
        nrPr.insert(0, nrFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        nrFonts.set(qn(attr), "Arial")
    normal.paragraph_format.space_after = Pt(BULLET_AFTER)
    normal.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    sec.header_distance = sec.footer_distance = Inches(0.3)

    # Header
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=0.5)
    _run(p, D.NAME, size=18, bold=True)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1.5)
    _run(p, title_line, size=10.5, bold=True, color=NAVY_HEX)
    for line in D.CONTACT_LINES:
        _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=0.5)
        _run(doc.paragraphs[-1], line)

    # Summary
    _section_heading(doc, "Summary")
    p = _para(doc, after=BLOCK_AFTER)
    _run(p, summary)

    # Skills
    _section_heading(doc, "Skills")
    for label, value in D.skills_for(skill_order):
        p = _para(doc, after=BULLET_AFTER, indent=BULLET_INDENT_PT,
                  hanging=BULLET_INDENT_PT)
        _run(p, "•  ")
        _run(p, label + ": ", bold=True)
        _run(p, value)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(BLOCK_AFTER)

    # Experience
    _section_heading(doc, "Experience")
    for i, job in enumerate(D.EXPERIENCE):
        p = _para(doc, after=JOB_AFTER, before=1.5 if i else 0.0)
        _run(p, job["header"], bold=True)
        for j, b in enumerate(job["bullets"]):
            last = j == len(job["bullets"]) - 1
            _bullet(doc, b, after=BLOCK_AFTER if (last and i == len(D.EXPERIENCE) - 1)
                    else BULLET_AFTER)

    # Projects
    _section_heading(doc, "Projects")
    for i, proj in enumerate(D.PROJECTS):
        p = _para(doc, after=JOB_AFTER, before=1.5 if i else 0.0)
        _run(p, proj["name"], bold=True)
        _run(p, " | " + proj["stack"])
        for j, b in enumerate(proj["bullets"]):
            last = j == len(proj["bullets"]) - 1
            _bullet(doc, b, after=BLOCK_AFTER if (last and i == len(D.PROJECTS) - 1)
                    else BULLET_AFTER)

    # Education
    _section_heading(doc, "Education")
    p = _para(doc, after=0.0)
    _run(p, D.EDUCATION)

    # Clean metadata: nothing that says how this file was made.
    cp = doc.core_properties
    cp.author = "Suthan K"
    cp.last_modified_by = "Suthan K"
    cp.title = "Suthan K - Resume"
    cp.subject = ""
    cp.comments = ""
    cp.category = ""
    cp.keywords = ""
    cp.revision = 1
    # python-docx's blank template carries a 2013 timestamp; make it today's.
    now = datetime.now().replace(microsecond=0)
    cp.created = now
    cp.modified = now

    doc.save(path)
    return path


# --------------------------------------------------------------------- the PDF
from reportlab import rl_config  # noqa: E402
from reportlab.lib.colors import Color, black  # noqa: E402
from reportlab.lib.enums import TA_CENTER, TA_LEFT  # noqa: E402
from reportlab.lib.styles import ParagraphStyle  # noqa: E402
from reportlab.pdfbase import pdfmetrics  # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont  # noqa: E402
from reportlab.platypus import (BaseDocTemplate, Frame, HRFlowable,  # noqa: E402
                                PageTemplate, Paragraph, Spacer)

NAVY = Color(*NAVY_RGB)
FONT, FONT_B = "Arial", "Arial-Bold"


def register_fonts():
    """Arial is not installed; Liberation Sans is metric-compatible with it."""
    pairs = [(FONT, ["LiberationSans-Regular.ttf", "Arial.ttf", "arial.ttf"]),
             (FONT_B, ["LiberationSans-Bold.ttf", "Arial_Bold.ttf", "arialbd.ttf"])]
    roots = ["/usr/share/fonts", "/usr/local/share/fonts",
             os.path.expanduser("~/.fonts"), os.path.expanduser("~/.local/share/fonts")]
    for name, candidates in pairs:
        for cand in candidates:
            hits = []
            for root in roots:
                hits += glob.glob(os.path.join(root, "**", cand), recursive=True)
            if hits:
                pdfmetrics.registerFont(TTFont(name, hits[0]))
                break
        else:
            raise SystemExit(f"no font file found for {name}")
    pdfmetrics.registerFontFamily(FONT, normal=FONT, bold=FONT_B,
                                  italic=FONT, boldItalic=FONT_B)
    # Stop ReportLab falling back to (non-embedded) Helvetica anywhere.
    rl_config.canvas_basefontname = FONT


def _styles():
    base = dict(fontName=FONT, fontSize=BODY_SIZE, leading=BODY_LEAD,
                textColor=black, spaceBefore=0, spaceAfter=BULLET_AFTER)
    s = {}
    s["name"] = ParagraphStyle("name", **{**base, "fontName": FONT_B,
                                          "fontSize": 18, "leading": 20.5,
                                          "alignment": TA_CENTER, "spaceAfter": 0.5})
    s["title"] = ParagraphStyle("title", **{**base, "fontName": FONT_B,
                                            "fontSize": 10.5, "leading": 12.5,
                                            "textColor": NAVY,
                                            "alignment": TA_CENTER, "spaceAfter": 1.5})
    s["contact"] = ParagraphStyle("contact", **{**base, "alignment": TA_CENTER,
                                                "spaceAfter": 0.5})
    s["section"] = ParagraphStyle("section", **{**base, "fontName": FONT_B,
                                                "fontSize": 11, "leading": 12.5,
                                                "alignment": TA_LEFT,
                                                "spaceBefore": SECTION_BEFORE,
                                                "spaceAfter": 0})
    s["body"] = ParagraphStyle("body", **{**base, "alignment": TA_LEFT})
    s["job"] = ParagraphStyle("job", **{**base, "fontName": FONT_B,
                                        "spaceBefore": 0, "spaceAfter": JOB_AFTER})
    s["bullet"] = ParagraphStyle("bullet", **{**base, "alignment": TA_LEFT,
                                              "leftIndent": BULLET_INDENT_PT + 4,
                                              "bulletIndent": 0,
                                              "bulletFontName": FONT,
                                              "bulletFontSize": BODY_SIZE})
    return s


def build_pdf(path, title_line, summary, skill_order):
    s = _styles()
    story = []

    def heading(text):
        story.append(Paragraph(text.upper(), s["section"]))
        story.append(HRFlowable(width="100%", thickness=0.75, color=NAVY,
                                spaceBefore=1.5, spaceAfter=SECTION_AFTER))

    def bullet(text, after=BULLET_AFTER):
        st = ParagraphStyle("b", parent=s["bullet"], spaceAfter=after)
        story.append(Paragraph(text, st, bulletText="•"))

    story.append(Paragraph(D.NAME, s["name"]))
    story.append(Paragraph(title_line, s["title"]))
    for line in D.CONTACT_LINES:
        story.append(Paragraph(line, s["contact"]))

    heading("Summary")
    story.append(Paragraph(summary, ParagraphStyle(
        "sum", parent=s["body"], spaceAfter=BLOCK_AFTER)))

    heading("Skills")
    skills = D.skills_for(skill_order)
    for i, (label, value) in enumerate(skills):
        after = BLOCK_AFTER if i == len(skills) - 1 else BULLET_AFTER
        bullet(f"<b>{label}:</b> {value}", after=after)

    heading("Experience")
    for i, job in enumerate(D.EXPERIENCE):
        story.append(Paragraph(job["header"], ParagraphStyle(
            "j", parent=s["job"], spaceBefore=1.5 if i else 0)))
        for j, b in enumerate(job["bullets"]):
            last = j == len(job["bullets"]) - 1 and i == len(D.EXPERIENCE) - 1
            bullet(b, after=BLOCK_AFTER if last else BULLET_AFTER)

    heading("Projects")
    for i, proj in enumerate(D.PROJECTS):
        story.append(Paragraph(
            f"<b>{proj['name']}</b> | {proj['stack']}",
            ParagraphStyle("p", parent=s["body"], fontName=FONT,
                           spaceBefore=1.5 if i else 0, spaceAfter=JOB_AFTER)))
        for j, b in enumerate(proj["bullets"]):
            last = j == len(proj["bullets"]) - 1 and i == len(D.PROJECTS) - 1
            bullet(b, after=BLOCK_AFTER if last else BULLET_AFTER)

    heading("Education")
    story.append(Paragraph(D.EDUCATION, ParagraphStyle(
        "edu", parent=s["body"], spaceAfter=0)))

    doc = BaseDocTemplate(
        path, pagesize=(A4_W_PT, A4_H_PT),
        leftMargin=MARGIN_PT, rightMargin=MARGIN_PT,
        topMargin=MARGIN_PT, bottomMargin=MARGIN_PT,
        title="Suthan K - Resume", author="Suthan K",
        subject="", creator="Suthan K", keywords="",
    )
    frame = Frame(MARGIN_PT, MARGIN_PT,
                  A4_W_PT - 2 * MARGIN_PT, A4_H_PT - 2 * MARGIN_PT,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
                  id="single")
    doc.addPageTemplates([PageTemplate(id="resume", frames=[frame])])
    doc.build(story)
    return path


# -------------------------------------------------------------------- the main
def main(argv):
    if not argv:
        raise SystemExit(__doc__)
    outdir, roles = argv[0], argv[1:] or list(D.ROLES)
    os.makedirs(outdir, exist_ok=True)
    register_fonts()
    for key, (suffix, title_line, summary, skill_order) in role_items(roles):
        dx = build_docx(os.path.join(outdir, out_name(suffix, "docx")),
                        title_line, summary, skill_order)
        pf = build_pdf(os.path.join(outdir, out_name(suffix, "pdf")),
                       title_line, summary, skill_order)
        print(f"{key:16s} {os.path.basename(dx):32s} {os.path.basename(pf)}")


if __name__ == "__main__":
    main(sys.argv[1:])
